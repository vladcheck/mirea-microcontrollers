#!/usr/bin/env python3
"""Restyle ЛР_1 report to match the accepted ПР_1-2 formatting conventions.

Conventions copied from pr1-2/ПР_1_Валекжанин_В_С.docx:
  - A4, margins L30/R15/T20/B20 mm (already correct in lab1)
  - title page in its own section WITHOUT footer; content section WITH
    centered PAGE-number footer (12 pt)
  - Normal style: Times New Roman 14, 1.5 spacing, NO style-level alignment
  - Heading 1: bold, color 365F91 (theme accent1, shade BF); centered via
    PER-PARAGRAPH jc, spacing before=12pt after=12pt; heading text stays
    CAPS + manual numbering ("1. ЦЕЛЬ РАБОТЫ")
  - body paragraphs: justified, first-line indent 709 twips (12.5 mm),
    spacing after=0
  - code listings: Consolas 10 pt, single spacing, default (left) alignment
  - figure/listing captions "Рисунок N — ..." / "Листинг N — ...": centered,
    spacing before=6pt after=6pt, TNR 14
  - image paragraphs: centered, spacing before=6pt
  - lists: plain Normal paragraphs (no auto-numbering), justified + 709 twips

Content is preserved 100%: no paragraph is added or deleted, no text is
changed (only formatting properties and the list numbering format are
touched). A timestamped backup is created next to the target before writing.

Usage:  uv run --with python-docx python restyle_report.py
"""

from __future__ import annotations

import shutil
import sys
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

BASE = Path(__file__).resolve().parent
TARGET = BASE / "ЛР_1_Валекжанин_В_С.docx"

# pr1-2 layout constants (twips)
FIRST_LINE = 709  # 12.51 mm first-line indent of body paragraphs
HEADING_SPACE = Pt(12)  # before/after for Heading 1 paragraphs
CAPTION_SPACE = Pt(6)  # before/after for captions, before for images
CODE_SIZE = Pt(10)
TITLE_PAGE_LAST_IDX = 23  # last title-page paragraph ("2026 г.")


def make_backup() -> Path:
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    bak = TARGET.parent / f"{TARGET.name}.bak-{ts}.docx"
    shutil.copy2(TARGET, bak)
    return bak


def set_run_fonts(run, name: str) -> None:
    """Set ascii/eastAsia/hAnsi/cs font of a run, dropping theme attrs."""
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:eastAsia", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:eastAsiaTheme", "w:hAnsiTheme", "w:cstheme"):
        if rf.get(qn(attr)) is not None:
            del rf.attrib[qn(attr)]


def remove_child(el, tag: str) -> None:
    for c in el.findall(qn(tag)):
        el.remove(c)


def fix_styles(doc: Document) -> None:
    # Normal: drop style-level justification (pr1-2 sets jc per paragraph)
    doc.styles["Normal"].paragraph_format.alignment = None

    # Heading 1: pr1-2 color 365F91 + theme accent1/BF, no style-level jc
    h1 = doc.styles["Heading 1"]
    h1.paragraph_format.alignment = None
    rpr = h1.element.get_or_add_rPr()
    remove_child(rpr, "w:color")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "365F91")
    color.set(qn("w:themeColor"), "accent1")
    color.set(qn("w:themeShade"), "BF")
    szcs = rpr.find(qn("w:szCs"))
    if szcs is not None:
        szcs.addprevious(color)
    else:
        rpr.append(color)


def fix_paragraphs(doc: Document) -> dict:
    stats = {
        "heading": 0,
        "image": 0,
        "caption": 0,
        "code": 0,
        "list": 0,
        "body": 0,
        "spacer": 0,
    }
    for idx, p in enumerate(doc.paragraphs):
        if idx <= TITLE_PAGE_LAST_IDX:
            continue  # title page already matches pr1-2 conventions
        pf = p.paragraph_format
        style = p.style.name
        text = p.text
        has_img = bool(p._p.findall(".//" + qn("a:blip")))
        is_caption = text.startswith(("Рисунок ", "Листинг "))
        is_code = any(r.font.size == CODE_SIZE for r in p.runs)

        if style == "Heading 1":
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = HEADING_SPACE
            pf.space_after = HEADING_SPACE
            pf.first_line_indent = None
            for r in p.runs:
                set_run_fonts(r, "Times New Roman")
            stats["heading"] += 1
        elif has_img:
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = CAPTION_SPACE
            pf.first_line_indent = None
            stats["image"] += 1
        elif is_caption:
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = CAPTION_SPACE
            pf.space_after = CAPTION_SPACE
            pf.first_line_indent = None
            stats["caption"] += 1
        elif is_code:
            pf.alignment = None  # default left, as in pr1-2
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pf.first_line_indent = None
            for r in p.runs:
                set_run_fonts(r, "Consolas")
            stats["code"] += 1
        elif style == "List Paragraph":  # auto-bullet list -> plain body
            p.style = doc.styles["Normal"]
            remove_child(p._p.get_or_add_pPr(), "w:numPr")
            pf.first_line_indent = Twips(FIRST_LINE)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_after = Pt(0)
            stats["list"] += 1
        elif text.strip() == "":
            pf.alignment = None  # spacer, as in pr1-2
            stats["spacer"] += 1
        else:  # body text incl. bold sub-headings
            pf.first_line_indent = Twips(FIRST_LINE)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_after = Pt(0)
            stats["body"] += 1
    return stats


def split_title_page_section(doc: Document) -> None:
    """Give the title page its own section without footer (like pr1-2)."""
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    title_sect = deepcopy(sect_pr)
    remove_child(title_sect, "w:footerReference")
    ppr = doc.paragraphs[TITLE_PAGE_LAST_IDX]._p.get_or_add_pPr()
    remove_child(ppr, "w:sectPr")
    ppr.append(title_sect)  # sectPr must be the last child of pPr


def snapshot(path: Path) -> dict:
    doc = Document(path)
    with zipfile.ZipFile(path) as z:
        media = sorted(n for n in z.namelist() if n.startswith("word/media/"))
    return {
        "texts": [p.text for p in doc.paragraphs],
        "media": media,
    }


def verify(bak: Path) -> bool:
    before = snapshot(bak)
    after_doc = Document(TARGET)
    with zipfile.ZipFile(TARGET) as z:
        bad = z.testzip()
        media_after = sorted(n for n in z.namelist() if n.startswith("word/media/"))
    after_texts = [p.text for p in after_doc.paragraphs]

    ok = True

    def check(label, cond, detail=""):
        nonlocal ok
        status = "OK " if cond else "FAIL"
        if not cond:
            ok = False
        print(f"  [{status}] {label}{(' — ' + detail) if detail else ''}")

    print("== VERIFICATION ==")
    check("zipfile.testzip()", bad is None, f"bad entry: {bad}")

    secs = after_doc.sections
    check(
        "section count == 2 (title page + content)", len(secs) == 2, f"got {len(secs)}"
    )
    for i, s in enumerate(secs):
        m = (
            s.left_margin.twips,
            s.right_margin.twips,
            s.top_margin.twips,
            s.bottom_margin.twips,
        )
        check(
            f"section {i} margins L/R/T/B twips", m == (1701, 850, 1134, 1134), str(m)
        )
        check(
            f"section {i} page size",
            (s.page_width.twips, s.page_height.twips) == (11906, 16838),
            f"{s.page_width.twips}x{s.page_height.twips}",
        )
    xml0 = secs[0]._sectPr
    xml1 = secs[1]._sectPr
    check(
        "section 0 (title page) has NO footer",
        not xml0.findall(qn("w:footerReference")),
    )
    check(
        "section 1 (content) has PAGE footer",
        bool(xml1.findall(qn("w:footerReference"))),
    )

    n = after_doc.styles["Normal"]
    check(
        "Normal font/size",
        n.font.name == "Times New Roman" and n.font.size == Pt(14),
        f"{n.font.name} {n.font.size}",
    )
    check("Normal line spacing 1.5", n.paragraph_format.line_spacing == 1.5)
    check("Normal style has no alignment (jc)", n.paragraph_format.alignment is None)

    h1 = after_doc.styles["Heading 1"]
    col = h1.element.get_or_add_rPr().find(qn("w:color"))
    check("Heading1 bold", h1.font.bold is True)
    check(
        "Heading1 color 365F91/accent1/BF",
        col is not None
        and col.get(qn("w:val")) == "365F91"
        and col.get(qn("w:themeColor")) == "accent1"
        and col.get(qn("w:themeShade")) == "BF",
    )
    check(
        "Heading1 style has no alignment (per-paragraph only)",
        h1.paragraph_format.alignment is None,
    )

    headings = [p for p in after_doc.paragraphs if p.style.name == "Heading 1"]
    print("  sample headings:")
    for p in headings[:3]:
        pf = p.paragraph_format
        print(
            f"    {p.text!r} | jc={pf.alignment} "
            f"sb={pf.space_before.pt if pf.space_before else None} "
            f"sa={pf.space_after.pt if pf.space_after else None}pt"
        )
    check(
        "all Heading1 centered, 12/12pt spacing",
        all(
            p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
            and p.paragraph_format.space_before == Pt(12)
            and p.paragraph_format.space_after == Pt(12)
            for p in headings
        ),
    )

    code_ps = [
        p for p in after_doc.paragraphs if any(r.font.size == Pt(10) for r in p.runs)
    ]
    fonts = {r.font.name for p in code_ps for r in p.runs}
    check(
        "code paragraphs use Consolas 10",
        fonts == {"Consolas"},
        f"fonts={fonts}, count={len(code_ps)}",
    )
    check(
        "code paragraphs have no explicit jc",
        all(p.paragraph_format.alignment is None for p in code_ps),
    )

    caps = [
        p for p in after_doc.paragraphs if p.text.startswith(("Рисунок ", "Листинг "))
    ]
    check(
        "all 5 listing captions + figure caption present",
        len(caps) == 6,
        f"found {len(caps)}",
    )
    check(
        "captions centered 6/6pt",
        all(
            p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
            and p.paragraph_format.space_before == Pt(6)
            and p.paragraph_format.space_after == Pt(6)
            for p in caps
        ),
    )
    for p in caps:
        print(f"    caption: {p.text!r}")

    listings = [p.text for p in caps if p.text.startswith("Листинг ")]
    check("5 listing captions", len(listings) == 5, f"{listings}")

    check(
        "paragraph count unchanged (zero add/delete)",
        len(before["texts"]) == len(after_texts),
        f"{len(before['texts'])} -> {len(after_texts)}",
    )
    diffs = sum(1 for a, b in zip(before["texts"], after_texts) if a != b)
    check("all paragraph texts identical to backup", diffs == 0, f"{diffs} text diffs")
    check(
        "image parts unchanged",
        before["media"] == media_after,
        f"{len(media_after)} parts",
    )
    check(
        "Cyrillic intact (spot checks)",
        any("ЦЕЛЬ РАБОТЫ" in t for t in after_texts)
        and any("Бегущий огонь" in t for t in after_texts)
        and any("Валекжанин" in t for t in after_texts),
    )
    return ok


def main() -> int:
    if not TARGET.exists():
        print(f"target not found: {TARGET}", file=sys.stderr)
        return 1
    bak = make_backup()
    print(f"backup: {bak}")

    doc = Document(TARGET)
    fix_styles(doc)
    stats = fix_paragraphs(doc)
    split_title_page_section(doc)
    doc.save(TARGET)
    print(f"restyled paragraphs: {stats}")

    return 0 if verify(bak) else 1


if __name__ == "__main__":
    sys.exit(main())
